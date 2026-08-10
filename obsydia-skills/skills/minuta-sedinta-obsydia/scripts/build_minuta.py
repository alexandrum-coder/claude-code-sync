#!/usr/bin/env python3
"""Generate an Obsydia meeting minute (.docx) from a JSON content file.

Usage:
    python3 build_minuta.py continut.json "Minuta_Subiect_10.08.2026.docx"

The document is built on assets/sablon-obsydia.docx, which carries the Obsydia
header and footer images in the Word section header/footer, so they repeat on
every page. The body of the template is cleared but its sectPr (margins, page
size, header/footer references) is preserved.

JSON schema (all keys required unless marked optional):

{
  "titlu":    "Minuta sedinta, Analiza Prima Pagina Portal <Client>",
  "proiect":  "Portal Digital <Client>, exemplu.ro",
  "data":     "10.08.2026",
  "ora":      "10:30, durata o ora",
  "locatie":  "Online (Microsoft Teams)",
  "client":   "<Client>",            # used only in the Participanti labels
  "obiectiv": ["...", "..."],
  "subiecte": [ {"titlu": "...", "puncte": ["...", "..."]} ],
  "decizii":  ["...", "..."],
  "actiuni":  [ ["actiune", "responsabil", "termen"] ],
  "pasi_intro": "...",               # optional, single line
  "pasi":     ["...", "..."]
}

Participants are never filled in from data: the two participant lines are always
emitted empty, to be completed manually in Pages.
"""

import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

NAVY = RGBColor(0x1A, 0x2C, 0x54)
NAVY_HEX = "1A2C54"
BORDER_HEX = "BFBFBF"
WIDTHS = (Cm(10.4), Cm(3.0), Cm(3.6))

FORBIDDEN = [
    "s-a solicitat", "s-au solicitat", "s-a cerut", "s-au cerut",
    "a cerut", "au cerut", "a solicitat", "au solicitat",
    "s-a stabilit", "s-a discutat", "s-a mentionat", "s-a menționat",
    "s-a propus", "a fost solicitat", "a fost cerut", "s-a convenit",
]
EMOJI = re.compile("[\U0001F300-\U0001FAFF←-⇿✀-➿✅⚠▪●]")


def build(spec, out_path, template):
    doc = Document(str(template))
    body = doc.element.body
    for el in list(body):
        if el.tag != qn("w:sectPr"):
            body.remove(el)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    def para(text="", bold=False, size=11, color=None, sb=0, sa=6,
             indent=0, align=None):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(sb)
        pf.space_after = Pt(sa)
        if indent:
            pf.left_indent = Cm(indent)
        if align is not None:
            p.alignment = align
        if text:
            r = p.add_run(text)
            r.bold = bold
            r.font.size = Pt(size)
            if color is not None:
                r.font.color.rgb = color
        return p

    def h1(text):
        return para(text, bold=True, size=13, color=NAVY, sb=16, sa=6)

    def labelled(label, value):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(label + " ")
        r.bold = True
        p.add_run(value)
        return p

    def bullet(text, indent=1.0):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Cm(indent)
        pf.first_line_indent = Cm(-0.5)
        pf.space_after = Pt(3)
        p.add_run("•  " + text)
        return p

    def numbered(i, text, indent=1.0):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Cm(indent)
        pf.first_line_indent = Cm(-0.75)
        pf.space_after = Pt(4)
        p.add_run("%d.  %s" % (i, text))
        return p

    def set_borders(table):
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement("w:" + edge)
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), "4")
            e.set(qn("w:space"), "0")
            e.set(qn("w:color"), BORDER_HEX)
            borders.append(e)
        table._tbl.tblPr.append(borders)

    def shade(cell, fill):
        sh = OxmlElement("w:shd")
        sh.set(qn("w:val"), "clear")
        sh.set(qn("w:fill"), fill)
        cell._tc.get_or_add_tcPr().append(sh)

    # --- title and identification block ---
    para(spec["titlu"], bold=True, size=16, color=NAVY, sa=14,
         align=WD_ALIGN_PARAGRAPH.CENTER)
    labelled("Proiect:", spec["proiect"])
    labelled("Data:", spec["data"])
    labelled("Ora:", spec["ora"])
    labelled("Locație:", spec["locatie"])

    # --- participants: always left empty, filled manually in Pages ---
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.add_run("Participanți:").bold = True
    client = spec.get("client", "").strip()
    bullet("Obsydia Software Solutions: ")
    bullet(("Reprezentanți %s: " % client).replace("  ", " "))

    # --- sections ---
    h1("Obiectivul ședinței")
    for b in spec["obiectiv"]:
        bullet(b)

    h1("Subiecte discutate")
    for i, s in enumerate(spec["subiecte"], 1):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(9)
        pf.space_after = Pt(2)
        pf.left_indent = Cm(0.5)
        pf.first_line_indent = Cm(-0.5)
        r = p.add_run("%d. %s" % (i, s["titlu"]))
        r.bold = True
        r.font.color.rgb = NAVY
        for b in s["puncte"]:
            bullet(b)

    h1("Decizii agreate")
    for i, d in enumerate(spec["decizii"], 1):
        numbered(i, d)

    h1("Acțiuni următoare")
    table = doc.add_table(rows=1, cols=3)
    set_borders(table)
    header = table.rows[0].cells
    for cell, text in zip(header, ("Acțiune", "Responsabil", "Termen")):
        shade(cell, NAVY_HEX)
        cell.paragraphs[0].paragraph_format.space_after = Pt(2)
        run = cell.paragraphs[0].add_run(text)
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        run.font.size = Pt(10)
    for row in spec["actiuni"]:
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            cell.paragraphs[0].paragraph_format.space_after = Pt(2)
            run = cell.paragraphs[0].add_run(text)
            run.font.size = Pt(10)
    for row in table.rows:
        for cell, width in zip(row.cells, WIDTHS):
            cell.width = width

    h1("Pași următori")
    if spec.get("pasi_intro"):
        para(spec["pasi_intro"], sa=4)
    for i, s in enumerate(spec["pasi"], 1):
        numbered(i, s)

    doc.save(str(out_path))
    return out_path


def verify(path, expected_rows):
    doc = Document(str(path))
    section = doc.sections[0]
    head = section.header._element.xml.count("<w:drawing>")
    foot = section.footer._element.xml.count("<w:drawing>")
    text = "\n".join(p.text for p in doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                text += "\n" + cell.text
    rows = len(doc.tables[0].rows) if doc.tables else 0

    problems = []
    print("antet: %d desen(e) | subsol: %d desen(e) | prima pagina diferita: %s"
          % (head, foot, section.different_first_page_header_footer))
    if head < 1 or foot < 1:
        problems.append("antet sau subsol lipsa")
    if section.different_first_page_header_footer:
        problems.append("prima pagina are antet diferit")

    print("randuri tabel: %d (asteptat %d)" % (rows, expected_rows))
    if rows != expected_rows:
        problems.append("numar de randuri neasteptat in tabel")

    dashes = text.count("—") + text.count("–")
    quotes = sum(text.count(c) for c in "„”“‟")
    emoji = len(EMOJI.findall(text))
    todo = text.count("[de verificat]")
    print("liniute lungi/medii: %d | ghilimele curbe: %d | emoji: %d"
          % (dashes, quotes, emoji))
    print("marcaje [de verificat]: %d" % todo)
    if dashes:
        problems.append("liniute lungi sau medii in text")
    if quotes:
        problems.append("ghilimele curbe in text")
    if emoji:
        problems.append("emoji sau iconite in text")

    low = text.lower()
    hits = [f for f in FORBIDDEN if f in low]
    if hits:
        print("formulari interzise: %s" % ", ".join(hits))
        problems.append("formulari interzise")
    else:
        print("formulari interzise: 0")

    if problems:
        print("PROBLEME: " + "; ".join(problems))
        return False
    print("OK")
    return True


def main():
    if len(sys.argv) < 3:
        print(__doc__)
        return 2
    spec_path = Path(sys.argv[1])
    out_path = Path(sys.argv[2])
    if out_path.exists():
        print("EROARE: %s exista deja. Alege alt nume." % out_path)
        return 1
    template = Path(__file__).resolve().parent.parent / "assets" / "sablon-obsydia.docx"
    if not template.exists():
        print("EROARE: sablonul lipseste: %s" % template)
        return 1
    spec = json.loads(spec_path.read_text(encoding="utf-8"))
    build(spec, out_path, template)
    print("Generat: %s" % out_path)
    return 0 if verify(out_path, len(spec["actiuni"]) + 1) else 1


if __name__ == "__main__":
    sys.exit(main())
